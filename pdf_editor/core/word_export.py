"""
Export a PDF to a Word (.docx) document.

Two backends are available:

* **layout** (:mod:`pdf2docx`) rebuilds paragraphs, tables, images and column
  layout, and is used whenever the package is installed.  This is the one that
  keeps the page looking like the original.
* **text** (:mod:`python-docx` plus PyMuPDF) is the fallback.  It walks the page
  content top to bottom and writes styled paragraphs and pictures.  Layout is
  approximate, but fonts, sizes, weights, colours and images survive.

Both accept an open document, so the GUI can export the edited state without
saving the PDF first.
"""

from __future__ import annotations

import io
import logging
import re
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, List, Optional, Sequence, Union

import pymupdf as fitz

from pdf_editor.core.fonts import SpanStyle
from pdf_editor.core.utils import PDFEditorError, UnsupportedOperationError

PathLike = Union[str, Path]
Source = Union[fitz.Document, bytes, PathLike]
ProgressCallback = Callable[[float, str], None]

# Points per inch, used to convert PDF geometry into Word measurements.
POINTS_PER_INCH = 72.0

# pdf2docx logs its progress as "(3/12) Page 3"; used to drive the progress bar.
_PROGRESS_PATTERN = re.compile(r"\((\d+)/(\d+)\)")


@dataclass
class WordExportReport:
    """Result of a Word export."""

    destination: Path
    pages: int
    method: str
    duration: float = 0.0
    warnings: List[str] = field(default_factory=list)

    @property
    def summary(self) -> str:
        """Short human readable description of the export."""
        style = "layout preserved" if self.method == "layout" else "text and images only"
        return (
            f"Exported {self.pages} page(s) to '{self.destination.name}' "
            f"({style}, {self.duration:.1f}s)"
        )


def layout_backend_available() -> bool:
    """
    Reports whether the layout-preserving backend can be used.

    :return: True when :mod:`pdf2docx` is importable.
    """
    try:
        import pdf2docx  # noqa: F401
    except Exception:
        return False
    return True


def text_backend_available() -> bool:
    """
    Reports whether the fallback backend can be used.

    :return: True when :mod:`python-docx` is importable.
    """
    try:
        import docx  # noqa: F401
    except Exception:
        return False
    return True


def available_methods() -> List[str]:
    """
    Lists the export methods this installation supports.

    :return: Any of ``"layout"`` and ``"text"``, best first.
    """
    methods = []
    if layout_backend_available():
        methods.append("layout")
    if text_backend_available():
        methods.append("text")
    return methods


def _document_bytes(source: Source) -> bytes:
    """
    Normalises any accepted source into PDF bytes.

    :param source: Open document, raw PDF bytes, or a path.
    :return: The PDF data.
    :raises PDFEditorError: When the source cannot be read.
    """
    if isinstance(source, fitz.Document):
        return source.tobytes(garbage=3, deflate=True)
    if isinstance(source, (bytes, bytearray)):
        return bytes(source)
    path = Path(source)
    if not path.exists():
        raise PDFEditorError(f"PDF file not found: '{path}'")
    try:
        return path.read_bytes()
    except OSError as error:
        raise PDFEditorError(f"Cannot read '{path.name}': {error}") from error


def _normalise_pages(pages: Optional[Sequence[int]], total: int) -> Optional[List[int]]:
    """
    Validates a page selection.

    :param pages: 0-based page indices, or ``None`` for the whole document.
    :param total: Number of pages in the document.
    :return: Sorted unique indices, or ``None`` for "all pages".
    :raises PDFEditorError: When an index is out of range.
    """
    if pages is None:
        return None
    selected = sorted({int(index) for index in pages})
    if not selected:
        raise PDFEditorError("No pages were selected for export.")
    for index in selected:
        if index < 0 or index >= total:
            raise PDFEditorError(f"Page {index + 1} does not exist (document has {total} page(s)).")
    return selected


class _ProgressLogHandler(logging.Handler):
    """Turns pdf2docx log lines into progress callbacks."""

    def __init__(self, callback: ProgressCallback) -> None:
        super().__init__(level=logging.INFO)
        self.callback = callback
        self.phase = 0
        self.last_step = 0

    def emit(self, record: logging.LogRecord) -> None:
        """Parses one log record and reports the progress it implies."""
        try:
            message = record.getMessage()
            match = _PROGRESS_PATTERN.search(message)
            if not match:
                return
            step, total = int(match.group(1)), int(match.group(2))
            if step < self.last_step:
                self.phase = 1  # pdf2docx runs two passes over the pages.
            self.last_step = step
            fraction = (self.phase + (step / max(total, 1))) / 2.0
            self.callback(min(0.99, fraction), f"Page {step} of {total}")
        except Exception:
            pass  # Progress reporting must never break an export.


def _export_with_layout(
    data: bytes,
    destination: Path,
    pages: Optional[List[int]],
    progress: Optional[ProgressCallback],
    quiet: bool = True,
) -> None:
    """
    Converts using :mod:`pdf2docx`, preserving the page layout.

    pdf2docx reports its progress by logging to the root logger, so the root
    logger is borrowed for the duration of the conversion: our handler reads the
    progress from it, and the existing console handlers are muted so the editor
    does not print conversion chatter.  Everything is restored afterwards.

    :param data: PDF bytes to convert.
    :param destination: Target .docx file.
    :param pages: 0-based page indices, or ``None`` for all pages.
    :param progress: Optional progress callback.
    :param quiet: Mute pdf2docx's own console output during the conversion.
    :raises PDFEditorError: When the conversion fails.
    """
    from pdf2docx import Converter  # Importing configures the root logger.

    root = logging.getLogger()
    handler = _ProgressLogHandler(progress) if progress else None
    previous_root_level = root.level
    muted: List[tuple] = []

    if handler:
        root.addHandler(handler)
        root.setLevel(min(previous_root_level or logging.INFO, logging.INFO))
    if quiet:
        for existing in list(root.handlers):
            if existing is not handler:
                muted.append((existing, existing.level))
                existing.setLevel(logging.WARNING)

    converter = None
    try:
        converter = Converter(stream=data)
        converter.convert(str(destination), pages=pages)
    except Exception as error:
        raise PDFEditorError(f"Word export failed: {error}") from error
    finally:
        if converter is not None:
            try:
                converter.close()
            except Exception:
                pass
        for existing, level in muted:
            existing.setLevel(level)
        if handler:
            root.removeHandler(handler)
            root.setLevel(previous_root_level)


def _style_run(run, style: SpanStyle) -> None:
    """
    Applies a PDF span style to a python-docx run.

    :param run: The docx run to format.
    :param style: Style read from the PDF.
    """
    from docx.shared import Pt, RGBColor

    run.font.size = Pt(max(1.0, style.size))
    run.font.bold = style.bold
    run.font.italic = style.italic
    red, green, blue = (int(round(component * 255)) for component in style.color)
    run.font.color.rgb = RGBColor(red, green, blue)
    family = "Courier New" if style.mono else ("Times New Roman" if style.serif else "Arial")
    run.font.name = family


def _export_with_text(
    data: bytes,
    destination: Path,
    pages: Optional[List[int]],
    progress: Optional[ProgressCallback],
) -> List[str]:
    """
    Converts with :mod:`python-docx`, writing styled text and the page images.

    :param data: PDF bytes to convert.
    :param destination: Target .docx file.
    :param pages: 0-based page indices, or ``None`` for all pages.
    :param progress: Optional progress callback.
    :return: Warnings collected during the export.
    :raises PDFEditorError: When the document cannot be written.
    """
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    warnings: List[str] = []
    document = fitz.open(stream=data, filetype="pdf")
    output = docx.Document()

    try:
        indices = pages if pages is not None else list(range(document.page_count))
        usable_width = Inches(6.5)

        for position, index in enumerate(indices):
            if progress:
                progress(position / max(len(indices), 1), f"Page {position + 1} of {len(indices)}")
            page = document[index]

            # Text blocks and images are interleaved in reading order.
            items = []
            for block in page.get_text("dict").get("blocks", []):
                if block.get("type", 1) == 0:
                    items.append((block["bbox"][1], block["bbox"][0], "text", block))
            for info in page.get_image_info(xrefs=True):
                if int(info.get("xref", 0)) > 0:
                    items.append((info["bbox"][1], info["bbox"][0], "image", info))
            items.sort(key=lambda item: (round(item[0], 1), item[1]))

            for _top, _left, kind, payload in items:
                if kind == "text":
                    _write_text_block(output, payload, page.rect.width, WD_ALIGN_PARAGRAPH)
                else:
                    written = _write_image(output, document, payload, usable_width)
                    if not written:
                        warnings.append(f"An image on page {index + 1} could not be exported.")

            if position < len(indices) - 1:
                output.add_page_break()

        destination.parent.mkdir(parents=True, exist_ok=True)
        output.save(str(destination))
    except PDFEditorError:
        raise
    except Exception as error:
        raise PDFEditorError(f"Word export failed: {error}") from error
    finally:
        document.close()
    return warnings


def _write_text_block(output, block: dict, page_width: float, alignments) -> None:
    """
    Writes one PDF text block as a Word paragraph.

    :param output: The docx document being built.
    :param block: PyMuPDF text block.
    :param page_width: Width of the page, used to guess the alignment.
    :param alignments: The ``WD_ALIGN_PARAGRAPH`` enumeration.
    """
    lines = block.get("lines", [])
    if not lines:
        return

    paragraph = output.add_paragraph()
    paragraph.paragraph_format.space_after = 0

    for line_index, line in enumerate(lines):
        for span in line.get("spans", []):
            text = span.get("text", "")
            if not text:
                continue
            run = paragraph.add_run(text)
            _style_run(run, SpanStyle.from_span(span))
        if line_index < len(lines) - 1:
            paragraph.add_run().add_break()

    left_gap = block["bbox"][0]
    right_gap = page_width - block["bbox"][2]
    if left_gap > 40 and abs(left_gap - right_gap) < 25:
        paragraph.alignment = alignments.CENTER
    elif left_gap > right_gap + 40:
        paragraph.alignment = alignments.RIGHT


def _write_image(output, document: fitz.Document, info: dict, usable_width) -> bool:
    """
    Writes one embedded image into the Word document.

    :param output: The docx document being built.
    :param document: Document holding the image.
    :param info: Entry from ``page.get_image_info(xrefs=True)``.
    :param usable_width: Maximum picture width.
    :return: True when the picture was added.
    """
    from docx.shared import Inches

    try:
        raw = document.extract_image(int(info["xref"]))
        stream = io.BytesIO(raw["image"])
        bbox = info["bbox"]
        width = Inches(min((bbox[2] - bbox[0]) / POINTS_PER_INCH, usable_width.inches))
        output.add_picture(stream, width=width)
        return True
    except Exception:
        return False


def export_to_docx(
    source: Source,
    destination: PathLike,
    *,
    pages: Optional[Sequence[int]] = None,
    method: str = "auto",
    progress: Optional[ProgressCallback] = None,
) -> WordExportReport:
    """
    Exports a PDF to a Word document.

    :param source: Open document, PDF bytes, or a path. Passing the open document
        exports the current state, including edits that were not saved yet.
    :param destination: Target ``.docx`` file; the suffix is added when missing.
    :param pages: 0-based page indices to export; ``None`` exports everything.
    :param method: ``"layout"``, ``"text"``, or ``"auto"`` to prefer layout.
    :param progress: Called as ``progress(fraction, message)`` during the export.
    :return: A :class:`WordExportReport`.
    :raises UnsupportedOperationError: When no backend is installed.
    :raises PDFEditorError: When the conversion or writing fails.
    """
    started = time.time()
    target = Path(destination)
    if target.suffix.lower() != ".docx":
        target = target.with_suffix(".docx")
    target.parent.mkdir(parents=True, exist_ok=True)

    data = _document_bytes(source)
    with fitz.open(stream=data, filetype="pdf") as probe:
        total_pages = probe.page_count
    selected = _normalise_pages(pages, total_pages)

    supported = available_methods()
    if not supported:
        raise UnsupportedOperationError(
            "Word export needs the 'pdf2docx' or 'python-docx' package. "
            "Install it with: pip install pdf2docx"
        )

    chosen = method if method != "auto" else supported[0]
    if chosen not in ("layout", "text"):
        raise PDFEditorError(f"Unknown export method '{method}'. Use 'layout', 'text' or 'auto'.")
    if chosen not in supported:
        raise UnsupportedOperationError(
            f"The '{chosen}' export method needs a package that is not installed "
            f"({'pdf2docx' if chosen == 'layout' else 'python-docx'})."
        )

    if progress:
        progress(0.0, "Preparing the document...")

    warnings: List[str] = []
    if chosen == "layout":
        _export_with_layout(data, target, selected, progress)
    else:
        warnings = _export_with_text(data, target, selected, progress)

    if not target.exists():
        raise PDFEditorError("The Word document was not created; the PDF may have no exportable content.")

    if progress:
        progress(1.0, "Finished")

    return WordExportReport(
        destination=target,
        pages=len(selected) if selected is not None else total_pages,
        method=chosen,
        duration=time.time() - started,
        warnings=warnings,
    )


__all__ = [
    "WordExportReport",
    "available_methods",
    "export_to_docx",
    "layout_backend_available",
    "text_backend_available",
]
