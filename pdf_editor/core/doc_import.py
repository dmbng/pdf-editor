"""
Import Word, Excel and PowerPoint documents and convert them to PDF.

Two backends are available and the best one is picked automatically:

* **office** drives an installed Microsoft Office through COM. Output is identical
  to "Save as PDF" in Word, Excel or PowerPoint, and it also reads the legacy
  ``.doc``, ``.xls`` and ``.ppt`` formats.
* **builtin** reads the file with ``python-docx`` / ``openpyxl`` / ``python-pptx``
  and draws the PDF with ReportLab. No Office needed, and it works on any
  platform, but the layout is a faithful approximation rather than a copy.

The converted PDF can be handed straight to the editor, so importing a document
and continuing to edit it is one step for the user.
"""

from __future__ import annotations

import gc
import html
import platform
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple, Union

from pdf_editor.core.utils import PDFEditorError, UnsupportedOperationError

PathLike = Union[str, Path]
ProgressCallback = Callable[[float, str], None]

# The three document families the editor can import.
KIND_WORD = "word"
KIND_EXCEL = "excel"
KIND_POWERPOINT = "powerpoint"

# Which extension belongs to which family.
EXTENSION_KINDS: Dict[str, str] = {
    ".docx": KIND_WORD,
    ".docm": KIND_WORD,
    ".doc": KIND_WORD,
    ".rtf": KIND_WORD,
    ".odt": KIND_WORD,
    ".xlsx": KIND_EXCEL,
    ".xlsm": KIND_EXCEL,
    ".xls": KIND_EXCEL,
    ".ods": KIND_EXCEL,
    ".csv": KIND_EXCEL,
    ".pptx": KIND_POWERPOINT,
    ".pptm": KIND_POWERPOINT,
    ".ppt": KIND_POWERPOINT,
    ".odp": KIND_POWERPOINT,
}

# Formats the pure Python backend can read on its own.
BUILTIN_EXTENSIONS = {".docx", ".docm", ".xlsx", ".xlsm", ".pptx", ".pptm"}

SUPPORTED_EXTENSIONS = set(EXTENSION_KINDS)

# COM identifiers and the "save as PDF" constants of each Office application.
_PROGIDS = {
    KIND_WORD: "Word.Application",
    KIND_EXCEL: "Excel.Application",
    KIND_POWERPOINT: "PowerPoint.Application",
}
_WORD_PDF_FORMAT = 17
_EXCEL_PDF_TYPE = 0
_POWERPOINT_PDF_FORMAT = 32

# Conversion factor from English Metric Units (used by Office) to PDF points.
EMU_PER_POINT = 12700.0


@dataclass
class ImportReport:
    """Result of a document conversion."""

    source: Path
    destination: Path
    kind: str
    backend: str
    pages: int = 0
    duration: float = 0.0
    warnings: List[str] = field(default_factory=list)

    @property
    def summary(self) -> str:
        """Short human readable description of the conversion."""
        engine = "Microsoft Office" if self.backend == "office" else "the built-in converter"
        return (
            f"Converted '{self.source.name}' to '{self.destination.name}' "
            f"({self.pages} page(s) via {engine}, {self.duration:.1f}s)"
        )


def detect_kind(source: PathLike) -> str:
    """
    Determines which application a file belongs to.

    :param source: File to inspect.
    :return: One of ``KIND_WORD``, ``KIND_EXCEL`` or ``KIND_POWERPOINT``.
    :raises UnsupportedOperationError: For a file type that cannot be imported.
    """
    suffix = Path(source).suffix.lower()
    kind = EXTENSION_KINDS.get(suffix)
    if kind is None:
        readable = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise UnsupportedOperationError(
            f"'{suffix or Path(source).name}' cannot be imported. Supported types: {readable}"
        )
    return kind


def is_supported(source: PathLike) -> bool:
    """
    Reports whether a file can be imported.

    :param source: File to inspect.
    :return: True when the extension is one the importer understands.
    """
    return Path(source).suffix.lower() in SUPPORTED_EXTENSIONS


def _com_client():
    """
    Returns a COM client module, or ``None`` when none is installed.

    :return: ``comtypes.client``, ``win32com.client``, or ``None``.
    """
    try:
        import comtypes.client  # noqa: F401

        return __import__("comtypes.client", fromlist=["client"])
    except Exception:
        pass
    try:
        import win32com.client  # noqa: F401

        return __import__("win32com.client", fromlist=["client"])
    except Exception:
        return None


def office_backend_available(kind: Optional[str] = None) -> bool:
    """
    Reports whether Microsoft Office can be driven on this machine.

    The check only reads the registry, so no Office window is ever opened.

    :param kind: Limit the check to one application, or ``None`` for any.
    :return: True when Office automation is usable.
    """
    if platform.system() != "Windows" or _com_client() is None:
        return False

    try:
        import winreg
    except Exception:
        return False

    wanted = [_PROGIDS[kind]] if kind else list(_PROGIDS.values())
    for prog_id in wanted:
        try:
            with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, prog_id):
                return True
        except OSError:
            continue
    return False


def builtin_backend_available(kind: Optional[str] = None) -> bool:
    """
    Reports whether the pure Python converter can run.

    :param kind: Limit the check to one application, or ``None`` for any.
    :return: True when the needed readers and ReportLab are installed.
    """
    try:
        import reportlab  # noqa: F401
    except Exception:
        return False

    readers = {
        KIND_WORD: "docx",
        KIND_EXCEL: "openpyxl",
        KIND_POWERPOINT: "pptx",
    }
    wanted = [readers[kind]] if kind else list(readers.values())
    for module in wanted:
        try:
            __import__(module)
            return True
        except Exception:
            continue
    return False


def available_backends(source: Optional[PathLike] = None) -> List[str]:
    """
    Lists the backends usable for a file, best first.

    :param source: File that would be converted, or ``None`` for a general answer.
    :return: Any of ``"office"`` and ``"builtin"``.
    """
    kind = detect_kind(source) if source is not None else None
    backends: List[str] = []
    if office_backend_available(kind):
        backends.append("office")

    if builtin_backend_available(kind):
        if source is None or Path(source).suffix.lower() in BUILTIN_EXTENSIONS:
            backends.append("builtin")
    return backends


def default_destination(source: PathLike) -> Path:
    """
    Returns the PDF path proposed for a source document.

    :param source: Document being converted.
    :return: The same folder and name with a ``.pdf`` suffix.
    """
    path = Path(source)
    return path.with_suffix(".pdf")


def convert_to_pdf(
    source: PathLike,
    destination: Optional[PathLike] = None,
    *,
    backend: str = "auto",
    progress: Optional[ProgressCallback] = None,
) -> ImportReport:
    """
    Converts a Word, Excel or PowerPoint document into a PDF file.

    :param source: Document to convert.
    :param destination: Target PDF; defaults to the source name with ``.pdf``.
    :param backend: ``"office"``, ``"builtin"`` or ``"auto"`` to prefer Office.
    :param progress: Called as ``progress(fraction, message)`` during the run.
    :return: An :class:`ImportReport`.
    :raises UnsupportedOperationError: When no usable backend is installed.
    :raises PDFEditorError: When the document cannot be read or written.
    """
    started = time.time()
    source_path = Path(source).resolve()
    if not source_path.exists() or not source_path.is_file():
        raise PDFEditorError(f"Document not found: '{source_path}'")

    kind = detect_kind(source_path)
    target = Path(destination) if destination else default_destination(source_path)
    if target.suffix.lower() != ".pdf":
        target = target.with_suffix(".pdf")
    target = target.resolve()
    if target == source_path:
        raise PDFEditorError("The PDF would overwrite the document being converted.")
    target.parent.mkdir(parents=True, exist_ok=True)

    usable = available_backends(source_path)
    if not usable:
        raise UnsupportedOperationError(_missing_backend_message(source_path, kind))

    chosen = backend if backend != "auto" else usable[0]
    if chosen not in ("office", "builtin"):
        raise PDFEditorError(f"Unknown backend '{backend}'. Use 'office', 'builtin' or 'auto'.")
    if chosen not in usable:
        raise UnsupportedOperationError(_missing_backend_message(source_path, kind, chosen))

    if progress:
        progress(0.05, f"Reading '{source_path.name}'...")

    warnings: List[str] = []
    if chosen == "office":
        _convert_with_office(source_path, target, kind, progress)
    else:
        warnings = _convert_with_builtin(source_path, target, kind, progress)

    if not target.exists():
        raise PDFEditorError("The converter did not produce a PDF file.")

    pages = _count_pages(target)
    if progress:
        progress(1.0, "Finished")

    return ImportReport(
        source=source_path,
        destination=target,
        kind=kind,
        backend=chosen,
        pages=pages,
        duration=time.time() - started,
        warnings=warnings,
    )


def _missing_backend_message(source: Path, kind: str, requested: Optional[str] = None) -> str:
    """
    Builds a helpful message when a conversion cannot be performed.

    :param source: Document the user tried to convert.
    :param kind: Application family of the document.
    :param requested: Backend that was asked for explicitly, if any.
    :return: A sentence naming what to install.
    """
    packages = {KIND_WORD: "python-docx", KIND_EXCEL: "openpyxl", KIND_POWERPOINT: "python-pptx"}
    if requested == "office":
        return (
            "Microsoft Office is not available on this computer, so the "
            "high fidelity converter cannot be used."
        )
    if requested == "builtin":
        if source.suffix.lower() not in BUILTIN_EXTENSIONS:
            return (
                f"The built-in converter cannot read '{source.suffix}' files. "
                "Only the modern formats (.docx, .xlsx, .pptx) are supported without Office."
            )
        return f"The built-in converter needs the '{packages[kind]}' and 'reportlab' packages."
    return (
        f"'{source.name}' cannot be converted here. Install Microsoft Office, "
        f"or install the '{packages[kind]}' and 'reportlab' packages for the built-in converter."
    )


def _count_pages(pdf_path: Path) -> int:
    """
    Counts the pages of a produced PDF.

    :param pdf_path: PDF to inspect.
    :return: Number of pages, or 0 when it cannot be opened.
    """
    try:
        import pymupdf as fitz

        with fitz.open(pdf_path) as document:
            return document.page_count
    except Exception:
        return 0


# ---------------------------------------------------------------------------
# Microsoft Office backend
# ---------------------------------------------------------------------------
class _ComSession:
    """Initialises COM for the calling thread and releases it afterwards."""

    def __init__(self) -> None:
        self._initialised = False

    def __enter__(self) -> "_ComSession":
        try:
            import pythoncom

            pythoncom.CoInitialize()
            self._initialised = True
        except Exception:
            self._initialised = False  # comtypes initialises COM on import.
        return self

    def __exit__(self, *exc_info) -> None:
        if not self._initialised:
            return
        # Every COM proxy must be released before COM itself is shut down,
        # otherwise the late release fails with an RPC error.
        gc.collect()
        try:
            import pythoncom

            pythoncom.CoUninitialize()
        except Exception:
            pass


def _dispatch(client, prog_id: str):
    """
    Creates a private instance of an Office application.

    :param client: The COM client module.
    :param prog_id: Application identifier such as ``"Word.Application"``.
    :return: The COM object.
    """
    if hasattr(client, "DispatchEx"):
        return client.DispatchEx(prog_id)  # A separate process, never a shared one.
    return client.CreateObject(prog_id)


def _convert_with_office(source: Path, target: Path, kind: str, progress) -> None:
    """
    Converts a document by driving Microsoft Office.

    :param source: Document to convert.
    :param target: PDF to write.
    :param kind: Application family of the document.
    :param progress: Optional progress callback.
    :raises PDFEditorError: When Office reports a problem.
    """
    client = _com_client()
    if client is None:
        raise UnsupportedOperationError("No COM client is installed (pywin32 or comtypes).")

    if progress:
        progress(0.3, "Starting Microsoft Office...")

    with _ComSession():
        # The conversion runs in its own frame so that every COM reference is
        # gone by the time the session shuts COM down.
        _drive_office(client, source, target, kind, progress)


def _drive_office(client, source: Path, target: Path, kind: str, progress) -> None:
    """
    Opens the document in Office and saves it as PDF.

    :param client: The COM client module.
    :param source: Document to convert.
    :param target: PDF to write.
    :param kind: Application family of the document.
    :param progress: Optional progress callback.
    :raises PDFEditorError: When Office reports a problem.
    """
    application = None
    try:
        application = _dispatch(client, _PROGIDS[kind])
        try:
            application.Visible = False
        except Exception:
            pass  # PowerPoint refuses to be hidden in some versions.
        try:
            application.DisplayAlerts = False
        except Exception:
            pass

        if progress:
            progress(0.6, "Converting...")

        # Every COM proxy is dropped as soon as it is closed: releasing one after
        # the application has quit fails with an RPC error.
        if kind == KIND_WORD:
            document = application.Documents.Open(
                str(source), ReadOnly=True, AddToRecentFiles=False, Visible=False
            )
            try:
                document.SaveAs(str(target), FileFormat=_WORD_PDF_FORMAT)
            finally:
                document.Close(False)
                document = None
        elif kind == KIND_EXCEL:
            workbook = application.Workbooks.Open(str(source), ReadOnly=True, UpdateLinks=0)
            try:
                workbook.ExportAsFixedFormat(_EXCEL_PDF_TYPE, str(target))
            finally:
                workbook.Close(False)
                workbook = None
        else:
            presentation = application.Presentations.Open(
                str(source), ReadOnly=True, WithWindow=False
            )
            try:
                presentation.SaveAs(str(target), _POWERPOINT_PDF_FORMAT)
            finally:
                presentation.Close()
                presentation = None
    except Exception as error:
        raise PDFEditorError(
            f"Microsoft Office could not convert '{source.name}': {error}"
        ) from error
    finally:
        if application is not None:
            try:
                application.Quit()
            except Exception:
                pass
            application = None
            gc.collect()


# ---------------------------------------------------------------------------
# Built-in backend
# ---------------------------------------------------------------------------
def _convert_with_builtin(source: Path, target: Path, kind: str, progress) -> List[str]:
    """
    Converts a document with the pure Python readers and ReportLab.

    :param source: Document to convert.
    :param target: PDF to write.
    :param kind: Application family of the document.
    :param progress: Optional progress callback.
    :return: Warnings collected during the conversion.
    :raises PDFEditorError: When the document cannot be read or drawn.
    """
    try:
        if kind == KIND_WORD:
            return _word_to_pdf(source, target, progress)
        if kind == KIND_EXCEL:
            return _excel_to_pdf(source, target, progress)
        return _powerpoint_to_pdf(source, target, progress)
    except PDFEditorError:
        raise
    except Exception as error:
        raise PDFEditorError(f"Could not convert '{source.name}': {error}") from error


def _escape(text: str) -> str:
    """
    Escapes text for ReportLab's mini HTML markup.

    :param text: Raw text.
    :return: Text safe to place inside a paragraph.
    """
    return html.escape(text or "", quote=False)


def _run_markup(run) -> str:
    """
    Renders one Word run as ReportLab markup.

    :param run: A ``python-docx`` run.
    :return: Marked up text.
    """
    text = _escape(run.text)
    if not text:
        return ""
    text = text.replace("\n", "<br/>").replace("\t", "&nbsp;&nbsp;&nbsp;&nbsp;")

    font = run.font
    colour = None
    try:
        if font.color is not None and font.color.rgb is not None:
            colour = f"#{str(font.color.rgb)}"
    except Exception:
        colour = None
    size = None
    try:
        if font.size is not None:
            size = font.size.pt
    except Exception:
        size = None

    if colour or size:
        attributes = ""
        if colour:
            attributes += f' color="{colour}"'
        if size:
            attributes += f' size="{size:.1f}"'
        text = f"<font{attributes}>{text}</font>"
    if run.bold:
        text = f"<b>{text}</b>"
    if run.italic:
        text = f"<i>{text}</i>"
    if run.underline:
        text = f"<u>{text}</u>"
    return text


def _paragraph_alignment(paragraph, alignments) -> int:
    """
    Maps a Word alignment onto a ReportLab one.

    :param paragraph: A ``python-docx`` paragraph.
    :param alignments: Module holding the ``TA_*`` constants.
    :return: The matching ReportLab alignment.
    """
    value = getattr(paragraph.paragraph_format, "alignment", None)
    name = str(value)
    if "CENTER" in name:
        return alignments.TA_CENTER
    if "RIGHT" in name:
        return alignments.TA_RIGHT
    if "JUSTIFY" in name:
        return alignments.TA_JUSTIFY
    return alignments.TA_LEFT


def _paragraph_images(paragraph, document) -> List[Tuple[bytes, float, float]]:
    """
    Extracts the pictures anchored in a Word paragraph.

    :param paragraph: A ``python-docx`` paragraph.
    :param document: The document the paragraph belongs to.
    :return: Tuples of ``(image bytes, width in points, height in points)``.
    """
    from docx.oxml.ns import qn

    pictures: List[Tuple[bytes, float, float]] = []
    for blip in paragraph._element.findall(f".//{qn('a:blip')}"):
        relationship_id = blip.get(qn("r:embed"))
        if not relationship_id:
            continue
        try:
            part = document.part.related_parts[relationship_id]
        except Exception:
            continue

        width = height = 0.0
        extent = blip.getparent().getparent().find(qn("wp:extent"))
        if extent is None:
            parent = blip.getparent().getparent().getparent()
            extent = parent.find(qn("wp:extent")) if parent is not None else None
        if extent is not None:
            try:
                width = float(extent.get("cx", 0)) / EMU_PER_POINT
                height = float(extent.get("cy", 0)) / EMU_PER_POINT
            except (TypeError, ValueError):
                width = height = 0.0
        pictures.append((part.blob, width, height))
    return pictures


def _word_to_pdf(source: Path, target: Path, progress) -> List[str]:
    """
    Renders a .docx file with ReportLab.

    :param source: Word document.
    :param target: PDF to write.
    :param progress: Optional progress callback.
    :return: Warnings collected during the conversion.
    """
    import io

    import docx
    from docx.oxml.ns import qn
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph
    from reportlab.lib import colors, enums
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import (
        Image,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    warnings: List[str] = []
    document = docx.Document(str(source))
    section = document.sections[0] if document.sections else None

    if section is not None:
        page_size = (section.page_width / EMU_PER_POINT, section.page_height / EMU_PER_POINT)
        margins = (
            section.left_margin / EMU_PER_POINT,
            section.right_margin / EMU_PER_POINT,
            section.top_margin / EMU_PER_POINT,
            section.bottom_margin / EMU_PER_POINT,
        )
    else:
        page_size = (595.0, 842.0)
        margins = (72.0, 72.0, 72.0, 72.0)

    template = SimpleDocTemplate(
        str(target),
        pagesize=page_size,
        leftMargin=margins[0],
        rightMargin=margins[1],
        topMargin=margins[2],
        bottomMargin=margins[3],
        title=source.stem,
    )
    available_width = page_size[0] - margins[0] - margins[1]

    styles = getSampleStyleSheet()
    body_style = ParagraphStyle(
        "ImportedBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=11,
        leading=14,
        spaceAfter=6,
    )

    def block_items():
        """Yields paragraphs and tables in document order."""
        for child in document.element.body.iterchildren():
            if child.tag == qn("w:p"):
                yield DocxParagraph(child, document)
            elif child.tag == qn("w:tbl"):
                yield DocxTable(child, document)

    story: List[object] = []
    items = list(block_items())

    for index, item in enumerate(items):
        if progress and index % 10 == 0:
            progress(0.1 + 0.8 * index / max(len(items), 1), f"Block {index + 1} of {len(items)}")

        if isinstance(item, DocxTable):
            story.append(_docx_table(item, available_width, body_style, Table, TableStyle, Paragraph, colors))
            story.append(Spacer(1, 8))
            continue

        style_name = (item.style.name if item.style is not None else "") or ""
        # A page break often sits in an otherwise empty paragraph, so it has to be
        # noticed before the "nothing to draw" shortcut below.
        page_break = bool(
            item._element.findall(f".//{qn('w:br')}[@{qn('w:type')}='page']")
        )

        for blob, width, height in _paragraph_images(item, document):
            try:
                if width <= 0 or height <= 0:
                    width, height = available_width * 0.6, available_width * 0.4
                scale = min(1.0, available_width / width) if width else 1.0
                story.append(Image(io.BytesIO(blob), width=width * scale, height=height * scale))
                story.append(Spacer(1, 6))
            except Exception:
                warnings.append("An image could not be converted.")

        markup = "".join(_run_markup(run) for run in item.runs)
        if not markup.strip():
            if item.text.strip():
                markup = _escape(item.text)
            else:
                story.append(Spacer(1, 6))
                if page_break:
                    story.append(PageBreak())
                continue

        if style_name.startswith("Heading"):
            level = "".join(character for character in style_name if character.isdigit())
            heading = styles.get(f"Heading{level or 1}", styles["Heading1"])
            paragraph_style = ParagraphStyle(
                f"Imported{style_name}", parent=heading, spaceBefore=10, spaceAfter=6
            )
        elif style_name.startswith("Title"):
            paragraph_style = ParagraphStyle("ImportedTitle", parent=styles["Title"])
        else:
            paragraph_style = ParagraphStyle(
                f"ImportedBody{index}",
                parent=body_style,
                alignment=_paragraph_alignment(item, enums),
                leftIndent=_indent(item, "left_indent"),
                firstLineIndent=_indent(item, "first_line_indent"),
            )

        bullet = None
        if "List Bullet" in style_name:
            bullet = "•"
        elif "List Number" in style_name:
            bullet = "-"
        if bullet:
            paragraph_style = ParagraphStyle(
                f"ImportedList{index}", parent=paragraph_style, leftIndent=18, bulletIndent=6
            )

        story.append(Paragraph(markup, paragraph_style, bulletText=bullet))
        if page_break:
            story.append(PageBreak())

    if not story:
        story.append(Paragraph(_escape(f"'{source.name}' contains no readable content."), body_style))

    if progress:
        progress(0.92, "Writing the PDF...")
    template.build(story)
    return warnings


def _indent(paragraph, attribute: str) -> float:
    """
    Reads a paragraph indent in points.

    :param paragraph: A ``python-docx`` paragraph.
    :param attribute: ``"left_indent"`` or ``"first_line_indent"``.
    :return: The indent in points, or 0.0 when unset.
    """
    try:
        value = getattr(paragraph.paragraph_format, attribute)
        return float(value.pt) if value is not None else 0.0
    except Exception:
        return 0.0


def _docx_table(table, available_width, body_style, Table, TableStyle, Paragraph, colors):
    """
    Converts a Word table into a ReportLab table.

    :param table: A ``python-docx`` table.
    :param available_width: Width the table may occupy.
    :param body_style: Base paragraph style for the cells.
    :param Table: The ReportLab ``Table`` class.
    :param TableStyle: The ReportLab ``TableStyle`` class.
    :param Paragraph: The ReportLab ``Paragraph`` class.
    :param colors: The ReportLab ``colors`` module.
    :return: A drawable table flowable.
    """
    from reportlab.lib.styles import ParagraphStyle

    cell_style = ParagraphStyle("ImportedCell", parent=body_style, fontSize=9, leading=11, spaceAfter=0)
    header_style = ParagraphStyle("ImportedCellHead", parent=cell_style, fontName="Helvetica-Bold")

    data = []
    for row_index, row in enumerate(table.rows):
        cells = []
        for cell in row.cells:
            text = "<br/>".join(_escape(paragraph.text) for paragraph in cell.paragraphs)
            cells.append(Paragraph(text or "&nbsp;", header_style if row_index == 0 else cell_style))
        data.append(cells)

    if not data:
        return Paragraph("", body_style)

    columns = max(len(row) for row in data)
    for row in data:
        while len(row) < columns:
            row.append(Paragraph("&nbsp;", cell_style))

    width = available_width / columns
    flowable = Table(data, colWidths=[width] * columns, repeatRows=1)
    flowable.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return flowable


def _format_cell(value) -> str:
    """
    Formats a spreadsheet value for display.

    :param value: Cell value from openpyxl.
    :return: The text to print.
    """
    import datetime as _datetime

    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (_datetime.datetime, _datetime.date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, _datetime.time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:g}"
    return str(value)


def _excel_to_pdf(source: Path, target: Path, progress) -> List[str]:
    """
    Renders an .xlsx workbook with ReportLab, one section per sheet.

    :param source: Excel workbook.
    :param target: PDF to write.
    :param progress: Optional progress callback.
    :return: Warnings collected during the conversion.
    """
    import openpyxl
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    warnings: List[str] = []
    workbook = openpyxl.load_workbook(str(source), data_only=True, read_only=True)

    page_size = landscape(A4)
    margin = 36.0
    template = SimpleDocTemplate(
        str(target),
        pagesize=page_size,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
        title=source.stem,
    )
    available_width = page_size[0] - 2 * margin

    styles = getSampleStyleSheet()
    heading_style = ParagraphStyle("SheetName", parent=styles["Heading2"], spaceAfter=8)
    cell_style = ParagraphStyle("Cell", parent=styles["BodyText"], fontSize=8, leading=10, spaceAfter=0)
    header_style = ParagraphStyle("CellHead", parent=cell_style, fontName="Helvetica-Bold")

    story: List[object] = []
    sheets = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
    if not sheets:
        sheets = list(workbook.worksheets)

    for sheet_index, sheet in enumerate(sheets):
        if progress:
            progress(
                0.1 + 0.8 * sheet_index / max(len(sheets), 1),
                f"Sheet '{sheet.title}' ({sheet_index + 1} of {len(sheets)})",
            )

        rows = [[_format_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
        while rows and not any(cell for cell in rows[-1]):
            rows.pop()
        if not rows:
            continue

        columns = max(len(row) for row in rows)
        last_used = 0
        for row in rows:
            for column_index, value in enumerate(row):
                if value:
                    last_used = max(last_used, column_index + 1)
        columns = max(1, min(columns, last_used or columns))

        data = []
        for row_index, row in enumerate(rows):
            padded = list(row[:columns]) + [""] * max(0, columns - len(row))
            style = header_style if row_index == 0 else cell_style
            data.append([Paragraph(_escape(value), style) for value in padded])

        if len(data) > 400:
            warnings.append(
                f"Sheet '{sheet.title}' has {len(data)} rows; the PDF will be long."
            )

        width = available_width / columns
        table = Table(data, colWidths=[width] * columns, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#bbbbbb")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef7")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 3),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ]
            )
        )

        story.append(Paragraph(_escape(sheet.title), heading_style))
        story.append(table)
        if sheet_index < len(sheets) - 1:
            story.append(PageBreak())

    workbook.close()
    if not story:
        story.append(Paragraph(_escape(f"'{source.name}' has no data to show."), cell_style))
        story.append(Spacer(1, 4))

    if progress:
        progress(0.92, "Writing the PDF...")
    template.build(story)
    return warnings


def _powerpoint_to_pdf(source: Path, target: Path, progress) -> List[str]:
    """
    Renders a .pptx presentation with ReportLab, one page per slide.

    :param source: PowerPoint presentation.
    :param target: PDF to write.
    :param progress: Optional progress callback.
    :return: Warnings collected during the conversion.
    """
    import io

    from pptx import Presentation
    from pptx.util import Emu
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas as pdfcanvas

    warnings: List[str] = []
    presentation = Presentation(str(source))
    width = float(presentation.slide_width) / EMU_PER_POINT
    height = float(presentation.slide_height) / EMU_PER_POINT

    surface = pdfcanvas.Canvas(str(target), pagesize=(width, height))
    surface.setTitle(source.stem)
    slides = list(presentation.slides)

    for index, slide in enumerate(slides):
        if progress:
            progress(0.1 + 0.8 * index / max(len(slides), 1), f"Slide {index + 1} of {len(slides)}")

        for shape in slide.shapes:
            try:
                _draw_shape(surface, shape, height, ImageReader, io, Emu)
            except Exception:
                warnings.append(f"A shape on slide {index + 1} could not be drawn.")
        surface.showPage()

    if not slides:
        surface.setFont("Helvetica", 14)
        surface.drawString(48, height - 60, f"'{source.name}' has no slides.")
        surface.showPage()

    if progress:
        progress(0.92, "Writing the PDF...")
    surface.save()
    return warnings


def _draw_shape(surface, shape, page_height: float, ImageReader, io, Emu) -> None:
    """
    Draws one PowerPoint shape onto the page.

    :param surface: The ReportLab canvas.
    :param shape: A ``python-pptx`` shape.
    :param page_height: Slide height in points, for the flipped Y axis.
    :param ImageReader: ReportLab's image reader class.
    :param io: The :mod:`io` module.
    :param Emu: ``pptx.util.Emu``.
    """
    left = float(shape.left or 0) / EMU_PER_POINT
    top = float(shape.top or 0) / EMU_PER_POINT
    box_width = float(shape.width or 0) / EMU_PER_POINT
    box_height = float(shape.height or 0) / EMU_PER_POINT

    if shape.shape_type is not None and str(shape.shape_type).startswith("PICTURE"):
        blob = shape.image.blob
        surface.drawImage(
            ImageReader(io.BytesIO(blob)),
            left,
            page_height - top - box_height,
            width=box_width,
            height=box_height,
            preserveAspectRatio=True,
            anchor="nw",
            mask="auto",
        )
        return

    if getattr(shape, "has_table", False) and shape.has_table:
        _draw_table(surface, shape, left, top, box_width, box_height, page_height)
        return

    if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
        return

    cursor = page_height - top
    for paragraph in shape.text_frame.paragraphs:
        text = "".join(run.text for run in paragraph.runs) or paragraph.text
        if not text.strip():
            cursor -= 12
            continue

        size = 18.0
        bold = italic = False
        colour = (0.0, 0.0, 0.0)
        for run in paragraph.runs:
            if run.font.size is not None:
                size = float(run.font.size.pt)
            bold = bool(run.font.bold) or bold
            italic = bool(run.font.italic) or italic
            try:
                if run.font.color is not None and run.font.color.rgb is not None:
                    red, green, blue = tuple(run.font.color.rgb)
                    colour = (red / 255.0, green / 255.0, blue / 255.0)
            except Exception:
                pass
            break

        name = "Helvetica"
        if bold and italic:
            name = "Helvetica-BoldOblique"
        elif bold:
            name = "Helvetica-Bold"
        elif italic:
            name = "Helvetica-Oblique"
        surface.setFont(name, size)
        surface.setFillColorRGB(*colour)

        alignment = str(paragraph.alignment or "")
        for line in _wrap_line(surface, text, name, size, box_width or 400):
            cursor -= size * 1.2
            if "CENTER" in alignment:
                surface.drawCentredString(left + (box_width or 0) / 2.0, cursor, line)
            elif "RIGHT" in alignment:
                surface.drawRightString(left + (box_width or 0), cursor, line)
            else:
                surface.drawString(left, cursor, line)
    surface.setFillColorRGB(0, 0, 0)


def _wrap_line(surface, text: str, font_name: str, size: float, width: float) -> List[str]:
    """
    Wraps a line of text to a given width.

    :param surface: The ReportLab canvas, used to measure text.
    :param text: Text to wrap.
    :param font_name: Font the text is drawn with.
    :param size: Font size in points.
    :param width: Available width in points.
    :return: The wrapped lines.
    """
    words = text.split()
    if not words:
        return []

    lines: List[str] = []
    current = words[0]
    for word in words[1:]:
        candidate = f"{current} {word}"
        if surface.stringWidth(candidate, font_name, size) <= max(width, 20):
            current = candidate
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


def _draw_table(surface, shape, left, top, box_width, box_height, page_height) -> None:
    """
    Draws a PowerPoint table as a simple grid.

    :param surface: The ReportLab canvas.
    :param shape: The shape holding the table.
    :param left: Left edge in points.
    :param top: Top edge in points.
    :param box_width: Table width in points.
    :param box_height: Table height in points.
    :param page_height: Slide height in points.
    """
    table = shape.table
    rows = len(table.rows)
    columns = len(table.columns)
    if not rows or not columns:
        return

    cell_width = box_width / columns
    cell_height = box_height / rows
    surface.setFont("Helvetica", min(11.0, cell_height * 0.5))
    surface.setStrokeColorRGB(0.6, 0.6, 0.6)

    for row_index in range(rows):
        for column_index in range(columns):
            x = left + column_index * cell_width
            y = page_height - top - (row_index + 1) * cell_height
            surface.rect(x, y, cell_width, cell_height, stroke=1, fill=0)
            text = table.cell(row_index, column_index).text.replace("\n", " ")
            if text:
                surface.drawString(x + 3, y + cell_height * 0.35, text[:120])


__all__ = [
    "BUILTIN_EXTENSIONS",
    "EXTENSION_KINDS",
    "ImportReport",
    "KIND_EXCEL",
    "KIND_POWERPOINT",
    "KIND_WORD",
    "SUPPORTED_EXTENSIONS",
    "available_backends",
    "builtin_backend_available",
    "convert_to_pdf",
    "default_destination",
    "detect_kind",
    "is_supported",
    "office_backend_available",
]