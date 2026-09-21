"""
Shared pytest fixtures for the PDF editor test suite.

The fixtures build their documents with PyMuPDF so the tests never depend on
files that happen to sit in the repository.
"""

import math
from pathlib import Path

import pymupdf as fitz
import pytest

BODY_TEXT = (
    "This paragraph exists so that the editing engine has something to reflow. "
    "It is deliberately long enough to wrap over several lines inside the text "
    "column, which is exactly the situation a replacement has to cope with."
)


@pytest.fixture
def paragraph_pdf(tmp_path: Path) -> Path:
    """
    A two page A4 document with a heading, a wrapped paragraph and a coloured line.

    :param tmp_path: Per-test temporary directory.
    :return: Path of the created PDF.
    """
    doc = fitz.open()

    page = doc.new_page(width=595, height=842)
    page.insert_text(fitz.Point(50, 80), "Editing Engine Test Page", fontsize=20, fontname="helv")
    page.insert_text(fitz.Point(50, 110), "Subtitle in red", fontsize=13, fontname="hebo", color=(0.8, 0.1, 0.1))
    page.insert_textbox(fitz.Rect(50, 140, 545, 260), BODY_TEXT, fontsize=11, fontname="helv")
    page.insert_textbox(
        fitz.Rect(50, 300, 545, 380),
        "A second paragraph follows the first one so that vertical growth is limited.",
        fontsize=11,
        fontname="helv",
    )

    second = doc.new_page(width=595, height=842)
    second.insert_text(fitz.Point(50, 80), "Second page heading", fontsize=16, fontname="helv")
    second.insert_text(fitz.Point(50, 120), "Reference number 12345", fontsize=11, fontname="cour")

    target = tmp_path / "paragraphs.pdf"
    doc.save(str(target))
    doc.close()
    return target


@pytest.fixture
def simple_pdf(tmp_path: Path) -> Path:
    """
    A four page document with one short line per page.

    :param tmp_path: Per-test temporary directory.
    :return: Path of the created PDF.
    """
    doc = fitz.open()
    for number in range(4):
        page = doc.new_page(width=400, height=500)
        page.insert_text(fitz.Point(40, 60), f"Page number {number + 1}", fontsize=14, fontname="helv")
    doc.set_metadata({"title": "Simple", "author": "Test Suite"})
    target = tmp_path / "simple.pdf"
    doc.save(str(target))
    doc.close()
    return target


@pytest.fixture
def other_pdf(tmp_path: Path) -> Path:
    """
    A two page document used as the source of merge operations.

    :param tmp_path: Per-test temporary directory.
    :return: Path of the created PDF.
    """
    doc = fitz.open()
    for number in range(2):
        page = doc.new_page(width=400, height=500)
        page.insert_text(fitz.Point(40, 60), f"Merged page {number + 1}", fontsize=14)
    doc.set_toc([[1, "Merged start", 1]])
    target = tmp_path / "other.pdf"
    doc.save(str(target))
    doc.close()
    return target


@pytest.fixture
def png_file(tmp_path: Path) -> Path:
    """
    A small RGB PNG used by the image tests.

    :param tmp_path: Per-test temporary directory.
    :return: Path of the created image.
    """
    pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 120, 60))
    pixmap.set_rect(pixmap.irect, (40, 120, 200))
    target = tmp_path / "picture.png"
    pixmap.save(str(target))
    return target


@pytest.fixture
def broken_pdf(tmp_path: Path) -> Path:
    """
    A file that is not a PDF at all.

    :param tmp_path: Per-test temporary directory.
    :return: Path of the created file.
    """
    target = tmp_path / "broken.pdf"
    target.write_bytes(b"this is definitely not a PDF file")
    return target


@pytest.fixture
def photo_pdf(tmp_path: Path) -> Path:
    """
    A two page document holding a high resolution photo-like image.

    The picture is placed small on the page, so its effective resolution is far
    above 200 dpi and every compression level has something to work on.

    :param tmp_path: Per-test temporary directory.
    :return: Path of the created PDF.
    """
    from PIL import Image, ImageDraw, ImageFilter

    width, height = 1200, 900
    image = Image.new("RGB", (width, height))
    pixels = image.load()
    for y in range(height):
        for x in range(width):
            pixels[x, y] = (
                int(120 + 100 * math.sin(x / 90.0)),
                int(110 + 90 * math.cos(y / 70.0)),
                int(140 + 80 * math.sin((x + y) / 110.0)),
            )
    draw = ImageDraw.Draw(image)
    for index in range(14):
        draw.ellipse(
            [index * 60, index * 40, index * 60 + 200, index * 40 + 160],
            outline=(250 - index * 10, 40 + index * 8, 90),
            width=7,
        )
    image = image.filter(ImageFilter.GaussianBlur(0.6))
    picture = tmp_path / "photo.jpg"
    image.save(picture, quality=92)

    doc = fitz.open()
    for number in range(2):
        page = doc.new_page(width=595, height=842)
        page.insert_text(fitz.Point(50, 60), f"Photo page {number + 1}", fontsize=16)
        # 1200 px across 240 pt is 360 dpi: well above every resampling threshold.
        page.insert_image(fitz.Rect(50, 90, 290, 270), filename=str(picture))
    target = tmp_path / "photos.pdf"
    doc.save(str(target), deflate=True)
    doc.close()
    return target


@pytest.fixture
def word_document(tmp_path: Path, png_file: Path) -> Path:
    """
    A .docx file with headings, styled runs, a bullet list, a table and a picture.

    :param tmp_path: Per-test temporary directory.
    :param png_file: Picture embedded in the document.
    :return: Path of the created document.
    """
    docx = pytest.importorskip("docx")
    from docx.shared import Inches, Pt, RGBColor

    document = docx.Document()
    document.add_heading("Imported Report", 0)
    document.add_heading("Section One", level=1)

    paragraph = document.add_paragraph("Plain text with ")
    run = paragraph.add_run("bold")
    run.bold = True
    paragraph.add_run(" and ")
    run = paragraph.add_run("red")
    run.font.color.rgb = RGBColor(0xC0, 0x20, 0x20)
    run.font.size = Pt(14)
    paragraph.add_run(" words.")

    document.add_paragraph("First bullet", style="List Bullet")
    document.add_paragraph("Second bullet", style="List Bullet")

    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue"
    table.cell(1, 0).text = "North"
    table.cell(1, 1).text = "120000"

    document.add_page_break()
    document.add_heading("Section Two", level=1)
    document.add_picture(str(png_file), width=Inches(2.0))
    document.add_paragraph("Closing line of the document.")

    target = tmp_path / "report.docx"
    document.save(str(target))
    return target


@pytest.fixture
def excel_workbook(tmp_path: Path) -> Path:
    """
    An .xlsx workbook with two sheets of values.

    :param tmp_path: Per-test temporary directory.
    :return: Path of the created workbook.
    """
    openpyxl = pytest.importorskip("openpyxl")

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["Region", "Q1", "Q2"])
    sheet.append(["North", 120000, 131000])
    sheet.append(["South", 98500, 97400])

    notes = workbook.create_sheet("Notes")
    notes.append(["Comment"])
    notes.append(["Figures are provisional."])

    target = tmp_path / "book.xlsx"
    workbook.save(str(target))
    return target


@pytest.fixture
def powerpoint_deck(tmp_path: Path, png_file: Path) -> Path:
    """
    A .pptx deck with a title slide and a slide holding a picture and a text box.

    :param tmp_path: Per-test temporary directory.
    :param png_file: Picture placed on the second slide.
    :return: Path of the created presentation.
    """
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    presentation = pptx.Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[0])
    first.shapes.title.text = "Deck Title"
    first.placeholders[1].text = "Subtitle line"

    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "Picture Slide"
    second.shapes.add_picture(str(png_file), Inches(1), Inches(2), width=Inches(3))
    box = second.shapes.add_textbox(Inches(5), Inches(2), Inches(3), Inches(2))
    box.text_frame.text = "Bullet one"
    box.text_frame.add_paragraph().text = "Bullet two"

    target = tmp_path / "deck.pptx"
    presentation.save(str(target))
    return target
